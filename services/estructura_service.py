from pathlib import Path

from openpyxl import load_workbook
from sqlalchemy import select

from database.conexion import obtener_sesion
from models.area import Area
from models.proyecto import Proyecto


class EstructuraService:

    def __init__(self):
        self.sesion = obtener_sesion()

    # ==================================================
    # LIMPIEZA
    # ==================================================

    @staticmethod
    def limpiar_texto(valor):
        if valor is None:
            return None

        valor = str(valor).strip()

        return valor if valor else None

    @staticmethod
    def limpiar_codigo(valor):
        if valor is None:
            return None

        # Por si Excel entrega 307.0
        if isinstance(valor, float) and valor.is_integer():
            return str(int(valor))

        return str(valor).strip()

    @staticmethod
    def limpiar_numero(valor):
        if valor is None:
            return None

        try:
            return int(float(valor))
        except (ValueError, TypeError):
            return None

    # ==================================================
    # ÁREAS
    # ==================================================

    def obtener_area(
        self,
        codigo_dg,
        codigo_da,
        descripcion
    ):
        consulta = select(Area).where(
            Area.codigo_dg == codigo_dg,
            Area.codigo_da == codigo_da,
            Area.descripcion == descripcion
        )

        return self.sesion.scalar(consulta)

    def obtener_o_crear_area(
        self,
        codigo_dg,
        codigo_da,
        descripcion,
        responsable,
        cargo
    ):
        area = self.obtener_area(
            codigo_dg,
            codigo_da,
            descripcion
        )

        if area:

            if responsable:
                area.responsable = responsable

            if cargo:
                area.cargo_responsable = cargo

            area.activo = True

            return area, False

        nueva_area = Area(
            codigo_dg=codigo_dg,
            codigo_da=codigo_da,
            descripcion=descripcion,
            responsable=responsable,
            cargo_responsable=cargo,
            activo=True
        )

        self.sesion.add(nueva_area)
        self.sesion.flush()

        return nueva_area, True

    # ==================================================
    # PROYECTOS
    # ==================================================

    def obtener_proyecto(self, codigo_proyecto):

        consulta = select(Proyecto).where(
            Proyecto.codigo_proyecto == codigo_proyecto
        )

        return self.sesion.scalar(consulta)

    # ==================================================
    # IMPORTACIÓN
    # ==================================================

    def importar_excel(self, ruta_archivo):

        ruta = Path(ruta_archivo)

        if not ruta.exists():
            raise FileNotFoundError(
                f"No se encontró el archivo: {ruta}"
            )

        # Abrir Excel directamente con openpyxl
        libro = load_workbook(
            filename=ruta,
            data_only=True
        )

        hoja = libro.active

        print()
        print(f"Hoja encontrada: {hoja.title}")
        print(f"Filas detectadas: {hoja.max_row}")
        print(f"Columnas detectadas: {hoja.max_column}")
        print()

        areas_creadas = 0
        areas_actualizadas = 0

        proyectos_creados = 0
        proyectos_actualizados = 0

        registros_omitidos = 0
        registros_leidos = 0

        try:

            # ==================================================
            # La fila 1 contiene encabezados.
            # Comenzamos desde la fila 2.
            #
            # B = No.
            # C = DG
            # D = DA
            # E = DESCRIPCION
            # F = Proyecto
            # G = Concepto
            # H = campo pendiente
            # I = campo pendiente
            # J = Nombre del Encargado
            # K = Area Responsable
            # ==================================================

            for fila_numero in range(2, hoja.max_row + 1):

                numero = hoja.cell(
                    row=fila_numero,
                    column=2
                ).value

                dg = hoja.cell(
                    row=fila_numero,
                    column=3
                ).value

                da = hoja.cell(
                    row=fila_numero,
                    column=4
                ).value

                descripcion = hoja.cell(
                    row=fila_numero,
                    column=5
                ).value

                proyecto = hoja.cell(
                    row=fila_numero,
                    column=6
                ).value

                concepto = hoja.cell(
                    row=fila_numero,
                    column=7
                ).value

                campo_h = hoja.cell(
                    row=fila_numero,
                    column=8
                ).value

                campo_i = hoja.cell(
                    row=fila_numero,
                    column=9
                ).value

                responsable = hoja.cell(
                    row=fila_numero,
                    column=10
                ).value

                cargo = hoja.cell(
                    row=fila_numero,
                    column=11
                ).value

                # ------------------------------------------
                # Si toda la fila está vacía, la ignoramos
                # ------------------------------------------

                valores = [
                    numero,
                    dg,
                    da,
                    descripcion,
                    proyecto,
                    concepto,
                    campo_h,
                    campo_i,
                    responsable,
                    cargo
                ]

                if all(valor is None for valor in valores):
                    continue

                registros_leidos += 1

                # ------------------------------------------
                # Limpieza
                # ------------------------------------------

                numero = self.limpiar_numero(numero)

                dg = self.limpiar_codigo(dg)

                da = self.limpiar_codigo(da)

                descripcion = self.limpiar_texto(
                    descripcion
                )

                codigo_proyecto = self.limpiar_codigo(
                    proyecto
                )

                concepto = self.limpiar_texto(
                    concepto
                )

                campo_h = self.limpiar_texto(
                    campo_h
                )

                campo_i = self.limpiar_texto(
                    campo_i
                )

                responsable = self.limpiar_texto(
                    responsable
                )

                cargo = self.limpiar_texto(
                    cargo
                )

                # ------------------------------------------
                # Validación
                # ------------------------------------------

                if not dg or not da or not descripcion:

                    registros_omitidos += 1

                    print(
                        f"Fila {fila_numero} omitida: "
                        f"faltan DG, DA o descripción."
                    )

                    continue

                # ------------------------------------------
                # Área
                # ------------------------------------------

                area, creada = self.obtener_o_crear_area(
                    codigo_dg=dg,
                    codigo_da=da,
                    descripcion=descripcion,
                    responsable=responsable,
                    cargo=cargo
                )

                if creada:
                    areas_creadas += 1
                else:
                    areas_actualizadas += 1

                # ------------------------------------------
                # Proyecto
                # ------------------------------------------

                if not codigo_proyecto or not concepto:

                    print(
                        f"Fila {fila_numero}: "
                        f"área importada sin proyecto."
                    )

                    continue

                proyecto_existente = (
                    self.obtener_proyecto(
                        codigo_proyecto
                    )
                )

                if proyecto_existente:

                    proyecto_existente.area_id = area.id
                    proyecto_existente.numero_origen = numero
                    proyecto_existente.concepto = concepto
                    proyecto_existente.campo_h = campo_h
                    proyecto_existente.campo_i = campo_i
                    proyecto_existente.activo = True

                    proyectos_actualizados += 1

                else:

                    nuevo_proyecto = Proyecto(
                        area_id=area.id,
                        numero_origen=numero,
                        codigo_proyecto=codigo_proyecto,
                        concepto=concepto,
                        campo_h=campo_h,
                        campo_i=campo_i,
                        activo=True
                    )

                    self.sesion.add(
                        nuevo_proyecto
                    )

                    proyectos_creados += 1

            # Guardamos sólo si terminó correctamente
            self.sesion.commit()

        except Exception:

            self.sesion.rollback()
            raise

        finally:

            libro.close()

        return {
            "registros_leidos": registros_leidos,

            "areas_creadas": areas_creadas,
            "areas_actualizadas": areas_actualizadas,

            "proyectos_creados": proyectos_creados,
            "proyectos_actualizados": proyectos_actualizados,

            "registros_omitidos": registros_omitidos
        }

    def cerrar(self):
        self.sesion.close()