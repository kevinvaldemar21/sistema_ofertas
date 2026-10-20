from pathlib import Path
from copy import copy

from openpyxl import load_workbook

import os


class ExcelService:

    def __init__(self):

        self.base_dir = (
            Path(__file__)
            .resolve()
            .parent
            .parent
        )

        self.templates_dir = (
            
            self.base_dir
            / "templates"
        )

        self.documentos_dir = (
            self.base_dir
            / "documentos"
        )

        self.documentos_dir.mkdir(
            exist_ok=True
        )

        self.plantilla = (
            self.templates_dir
            / "ESTUDIO_DE_MERCADO.xlsx"
        )

    # ==================================================
    # FORMATO DE FECHA
    # ==================================================

    @staticmethod
    def formatear_fecha(fecha):

        meses = [
            "",
            "enero",
            "febrero",
            "marzo",
            "abril",
            "mayo",
            "junio",
            "julio",
            "agosto",
            "septiembre",
            "octubre",
            "noviembre",
            "diciembre"
        ]

        return (
            f"{fecha.day:02d} de "
            f"{meses[fecha.month]} "
            f"{fecha.year}"
        )


    # ==================================================
    # NOMBRE DEL DOCUMENTO
    # ==================================================

    @staticmethod
    def nombre_documento(estudio):

        expediente_seguro = (
            estudio.expediente
            .replace("/", "-")
            .replace("\\", "-")
            .replace(":", "-")
        )

        return (
            f"ESTUDIO_DE_MERCADO_"
            f"{expediente_seguro}.xlsx"
        )


    # ==================================================
    # RUTA DEL DOCUMENTO
    # ==================================================

    def obtener_ruta_documento(
        self,
        estudio
    ):

        return (
            self.documentos_dir
            / self.nombre_documento(
                estudio
            )
        )


    # ==================================================
    # VERIFICAR SI EXISTE
    # ==================================================

    def documento_existe(
        self,
        estudio
    ):

        ruta = (
            self.obtener_ruta_documento(
                estudio
            )
        )

        return ruta.exists()

    # ==================================================
    # GENERAR ESTUDIO
    # ==================================================

    def generar_estudio(
        self,
        estudio,
        cotizaciones,
        cotizacion_repository, 
        condiciones,
        evaluaciones
    ):

        if not self.plantilla.exists():

            raise FileNotFoundError(
                "No se encontró la plantilla:\n"
                f"{self.plantilla}"
            )

        libro = load_workbook(
            self.plantilla,
            data_only=False,
            keep_links=False
        )

        self.eliminar_referencias_externas(
            libro
        )

        hoja = libro["Hoja1"]

        # ==============================================
        # DATOS GENERALES
        # ==============================================

        hoja["K3"] = (
            self.formatear_fecha(
                estudio.fecha_elaboracion
            )
        )

        hoja["K4"] = (
            estudio.expediente
        )


        # ==================================================
        # PRESTADORES CONSULTADOS
        # ==================================================

        celdas_prestadores = [
            {
                "nombre": "C16",
                "telefono": "G16",
                "padron": "J16"
            },
            {
                "nombre": "C17",
                "telefono": "G17",
                "padron": "J17"
            },
            {
                "nombre": "C18",
                "telefono": "G18",
                "padron": "J18"
            }
        ]


        for indice in range(3):

            if indice < len(cotizaciones):

                prestador = (
                    cotizaciones[indice]
                    .prestador
                )

                hoja[
                    celdas_prestadores[indice]["nombre"]
                ] = (
                    prestador.nombre_razon_social
                )

                hoja[
                    celdas_prestadores[indice]["telefono"]
                ] = (
                    prestador.telefono or ""
                )

                hoja[
                    celdas_prestadores[indice]["padron"]
                ] = (
                    prestador.registro_padron or ""
                )

            else:

                # Si sólo participaron dos prestadores,
                # limpiamos la tercera columna de la plantilla.

                hoja[
                    celdas_prestadores[indice]["nombre"]
                ] = ""

                hoja[
                    celdas_prestadores[indice]["telefono"]
                ] = ""

                hoja[
                    celdas_prestadores[indice]["padron"]
                ] = ""

        
        # ==================================================
        # FECHA DEL COMPARATIVO
        # ==================================================

        hoja["D43"] = (
            estudio.fecha_elaboracion
        )

        hoja["D43"].number_format = (
            "dd/mm/yyyy"
        )


        # ==================================================
        # ENCABEZADOS DE COTIZACIONES
        # ==================================================

        celdas_cotizaciones = [
            {
                "nombre": "F43",
                "rfc": "F44"
            },
            {
                "nombre": "H43",
                "rfc": "H44"
            },
            {
                "nombre": "J43",
                "rfc": "J44"
            }
        ]

        for indice in range(3):

            if indice < len(cotizaciones):

                prestador = (
                    cotizaciones[indice]
                    .prestador
                )

                hoja[
                    celdas_cotizaciones[indice]["nombre"]
                ] = (
                    prestador.nombre_razon_social
                )

                hoja[
                    celdas_cotizaciones[indice]["rfc"]
                ] = (
                    prestador.rfc or ""
                )

            else:

                hoja[
                    celdas_cotizaciones[indice]["nombre"]
                ] = None

                hoja[
                    celdas_cotizaciones[indice]["rfc"]
                ] = None


        # ==================================================
        # TABLA PRINCIPAL DE CONCEPTOS Y COTIZACIONES
        # ==================================================

        fila_inicio = 46

        detalles_por_cotizacion = []

        for cotizacion in cotizaciones:

            detalles = (
                cotizacion_repository
                .obtener_detalles(
                    cotizacion.id
                )
            )

            detalles_por_cotizacion.append(
                detalles
            )


        # Usamos la primera cotización para obtener
        # los conceptos del estudio
        if detalles_por_cotizacion:

            detalles_base = (
                detalles_por_cotizacion[0]
            )

        else:

            detalles_base = []


        for indice, detalle_base in enumerate(
            detalles_base
        ):

            fila = fila_inicio + indice

            concepto = detalle_base.concepto

            # ----------------------------------------------
            # DATOS DEL CONCEPTO
            # ----------------------------------------------

            hoja[f"B{fila}"] = (
                concepto.numero
            )

            hoja[f"C{fila}"] = (
                concepto.descripcion
            )

            hoja[f"D{fila}"] = (
                concepto.unidad_medida
            )

            hoja[f"E{fila}"] = (
                float(
                    concepto.cantidad
                )
            )

            # ----------------------------------------------
            # DATOS DE CADA PRESTADOR
            # ----------------------------------------------

            columnas = [
                ("F", "G"),
                ("H", "I"),
                ("J", "K")
            ]

            precios_validos = []

            for indice_prestador in range(3):

                columna_precio, columna_importe = (
                    columnas[indice_prestador]
                )

                if (
                    indice_prestador
                    < len(detalles_por_cotizacion)
                ):

                    detalles = (
                        detalles_por_cotizacion[
                            indice_prestador
                        ]
                    )

                    detalle = detalles[indice]

                    precio = float(
                        detalle.precio_unitario
                    )

                    importe = float(
                        detalle.importe
                    )

                    hoja[
                        f"{columna_precio}{fila}"
                    ] = precio

                    hoja[
                        f"{columna_importe}{fila}"
                    ] = importe

                    precios_validos.append(
                        precio
                    )

                else:

                    hoja[
                        f"{columna_precio}{fila}"
                    ] = ""

                    hoja[
                        f"{columna_importe}{fila}"
                    ] = ""

            # ----------------------------------------------
            # PRECIO PROMEDIO DE REFERENCIA
            # ----------------------------------------------

            if precios_validos:

                promedio = (
                    sum(precios_validos)
                    / len(precios_validos)
                )

                hoja[f"L{fila}"] = round(
                    promedio,
                    2
                )

            else:

                hoja[f"L{fila}"] = ""


        # ==================================================
        # SUBTOTAL, IVA Y TOTAL
        # ==================================================

        columnas_totales = [
            "G",
            "I",
            "K"
        ]


        for indice in range(3):

            if indice < len(cotizaciones):

                cotizacion = (
                    cotizaciones[indice]
                )

                columna = (
                    columnas_totales[indice]
                )

                hoja[
                    f"{columna}50"
                ] = float(
                    cotizacion.subtotal
                )

                hoja[
                    f"{columna}51"
                ] = float(
                    cotizacion.iva
                )

                hoja[
                    f"{columna}52"
                ] = float(
                    cotizacion.total
                )

            else:

                columna = (
                    columnas_totales[indice]
                )

                hoja[
                    f"{columna}50"
                ] = ""

                hoja[
                    f"{columna}51"
                ] = ""

                hoja[
                    f"{columna}52"
                ] = ""


                # ==================================================
        # PROMEDIO GENERAL DE REFERENCIA
        # ==================================================

        totales_validos = [
            float(cotizacion.total)
            for cotizacion in cotizaciones
        ]

        if totales_validos:

            hoja["L52"] = round(
                sum(totales_validos)
                / len(totales_validos),
                2
            )

        else:

            hoja["L52"] = ""


                # ==================================================
        # FORMATO MONETARIO
        # ==================================================

        formato_moneda = '$#,##0.00'


        for fila in range(
            fila_inicio,
            fila_inicio + len(detalles_base)
        ):

            for columna in [
                "F",
                "G",
                "H",
                "I",
                "J",
                "K",
                "L"
            ]:

                hoja[
                    f"{columna}{fila}"
                ].number_format = formato_moneda


        for columna in [
            "G",
            "I",
            "K",
            "L"
        ]:

            for fila in [
                50,
                51,
                52
            ]:

                hoja[
                    f"{columna}{fila}"
                ].number_format = formato_moneda
        
        # ==================================================
        # CANTIDAD, CALIDAD Y OPORTUNIDAD
        # ==================================================

        filas_cantidad = [
            24,
            25,
            26
        ]

        filas_calidad = [
            30,
            31,
            32
        ]

        filas_oportunidad = [
            36,
            37,
            38
        ]


        # Creamos un diccionario para localizar rápidamente
        # la evaluación de cada prestador.
        evaluaciones_por_prestador = {
            evaluacion.prestador_id: evaluacion
            for evaluacion in evaluaciones
        }


        for indice in range(3):

            if indice < len(cotizaciones):

                prestador = (
                    cotizaciones[indice]
                    .prestador
                )

                evaluacion = (
                    evaluaciones_por_prestador
                    .get(
                        prestador.id
                    )
                )

                # ==========================================
                # CANTIDAD
                # ==========================================

                fila = (
                    filas_cantidad[indice]
                )

                hoja[f"C{fila}"] = (
                    prestador.nombre_razon_social
                )

                hoja[f"G{fila}"] = (
                    prestador.telefono or ""
                )

                if evaluacion:

                    hoja[f"J{fila}"] = (
                        "SI"
                        if evaluacion.cumple_cantidad
                        else "NO"
                    )

                else:

                    hoja[f"J{fila}"] = ""

                # ==========================================
                # CALIDAD
                # ==========================================

                fila = (
                    filas_calidad[indice]
                )

                hoja[f"C{fila}"] = (
                    prestador.nombre_razon_social
                )

                hoja[f"G{fila}"] = (
                    prestador.telefono or ""
                )

                if evaluacion:

                    hoja[f"J{fila}"] = (
                        "SI"
                        if evaluacion.cumple_calidad
                        else "NO"
                    )

                else:

                    hoja[f"J{fila}"] = ""

                # ==========================================
                # OPORTUNIDAD
                # ==========================================

                fila = (
                    filas_oportunidad[indice]
                )

                hoja[f"C{fila}"] = (
                    prestador.nombre_razon_social
                )

                hoja[f"G{fila}"] = (
                    prestador.telefono or ""
                )

                if evaluacion:

                    hoja[f"J{fila}"] = (
                        "SI"
                        if evaluacion.cumple_oportunidad
                        else "NO"
                    )

                else:

                    hoja[f"J{fila}"] = ""

            else:

                # ==========================================
                # LIMPIAR ESPACIOS NO UTILIZADOS
                # ==========================================

                for fila in [
                    filas_cantidad[indice],
                    filas_calidad[indice],
                    filas_oportunidad[indice]
                ]:

                    hoja[f"C{fila}"] = ""
                    hoja[f"G{fila}"] = ""
                    hoja[f"J{fila}"] = ""

        
        

        # ==================================================
        # PRESTADOR SELECCIONADO
        # ==================================================

        cotizacion_seleccionada = None
        indice_seleccionado = None

        for indice, cotizacion in enumerate(
            cotizaciones
        ):

            if cotizacion.seleccionada:

                cotizacion_seleccionada = (
                    cotizacion
                )

                indice_seleccionado = (
                    indice
                )

                break


        if cotizacion_seleccionada:

            prestador_seleccionado = (
                cotizacion_seleccionada
                .prestador
            )

            nombre_prestador = (
                prestador_seleccionado
                .nombre_razon_social
            )

            rfc_prestador = (
                prestador_seleccionado
                .rfc or ""
            )

            domicilio_prestador = (
                prestador_seleccionado
                .domicilio or ""
            )

            total_seleccionado = float(
                cotizacion_seleccionada.total
            )

        else:

            nombre_prestador = ""
            rfc_prestador = ""
            domicilio_prestador = ""
            total_seleccionado = 0

        

        # ==================================================
        # TOTAL DEL PRESTADOR SELECCIONADO
        # ==================================================

        celdas_totales_prestadores = [
            "G52",
            "I52",
            "K52"
        ]


        if (
            cotizacion_seleccionada is not None
            and indice_seleccionado is not None
        ):

            celda_total_seleccionado = (
                celdas_totales_prestadores[
                    indice_seleccionado
                ]
            )

            total_seleccionado_letra = (
                self.numero_a_letras_mxn(
                    total_seleccionado
                )
            )

            # Total numérico
            hoja["H61"] = (
                f"={celda_total_seleccionado}"
            )

            hoja["H61"].number_format = (
                '$#,##0.00'
            )

            # Total en letra
            hoja["I61"] = (
                total_seleccionado_letra
            )


        # ==================================================
        # CONDICIONES COMERCIALES
        # ==================================================

        if condiciones:

            hoja["C56"] = (
                condiciones.lugar_entrega
            )

            hoja["C57"] = (
                condiciones.plazo_entrega
            )

            hoja["C58"] = (
                condiciones.forma_pago
            )

        # ==================================================
        # TIPO DE PROCEDIMIENTO
        # ==================================================

        if condiciones and condiciones.tipo_procedimiento:

            hoja["B72"] = (
                "TIPO DE PROCEDIMIENTO: "
                + condiciones.tipo_procedimiento.upper()
            )

        # ==================================================
        # DATOS DEL PRESTADOR SELECCIONADO
        # ==================================================

        if cotizacion_seleccionada:

            hoja["C66"] = (
                nombre_prestador
            )

            hoja["G66"] = (
                rfc_prestador
            )

            hoja["I66"] = (
                domicilio_prestador
            )


        # ==================================================
        # PRECIO TOTAL OFERTADO
        # ==================================================

        hoja["E69"] = "=H61"
        hoja["F69"] = "=I61"

        hoja["E69"].number_format = (
            '$#,##0.00'
        )


        # ==================================================
        # DETERMINACIÓN FINAL
        # ==================================================

        hoja["C75"] = "=C66"
        hoja["G75"] = "=G66"
        hoja["I75"] = "=I66"


        # ==================================================
        # PRECIO TOTAL CON IMPUESTOS
        # ==================================================

        hoja["E76"] = "=E69"
        hoja["F76"] = "=F69"

        hoja["E76"].number_format = (
            '$#,##0.00'
        )
        
        # ==============================================
        # INFORMACIÓN PROGRAMÁTICA
        # ==============================================

        area_completa = (
            f"{estudio.codigo_dg} "
            f"{estudio.codigo_da} "
            f"{estudio.area_descripcion}"
        )

        hoja["D9"] = (
            area_completa
        )

        hoja["K9"] = (
            estudio.codigo_proyecto
        )

        hoja["D10"] = (
            estudio.partida_presupuestal
        )

        hoja["D11"] = (
            estudio.objeto_contratacion
        )

        # ==============================================
        # INFORMACIÓN REPETIDA AL FINAL
        # ==============================================

        hoja["D77"] = (
            area_completa
        )

        hoja["L77"] = (
            estudio.codigo_proyecto
        )

        hoja["D78"] = (
            estudio.partida_presupuestal
        )

        # ==================================================
        # GUARDADO SEGURO DEL DOCUMENTO
        # ==================================================

        ruta_salida = (
            self.obtener_ruta_documento(
                estudio
            )
        )

        ruta_temporal = (
            self.documentos_dir
            / (
                ruta_salida.stem
                + "_TEMPORAL.xlsx"
            )
        )

        try:

            # Guardar primero una copia temporal
            libro.save(
                ruta_temporal
            )

            libro.close()

            # Si todo salió correctamente,
            # sustituir el archivo anterior
            os.replace(
                ruta_temporal,
                ruta_salida
            )

        except Exception:

            try:

                libro.close()

            except Exception:

                pass

            if ruta_temporal.exists():

                try:

                    ruta_temporal.unlink()

                except Exception:

                    pass

            raise


        return ruta_salida

    @staticmethod   
    def numero_a_letras_mxn(numero):

        unidades = [
            "",
            "UNO",
            "DOS",
            "TRES",
            "CUATRO",
            "CINCO",
            "SEIS",
            "SIETE",
            "OCHO",
            "NUEVE"
        ]

        especiales = {
            10: "DIEZ",
            11: "ONCE",
            12: "DOCE",
            13: "TRECE",
            14: "CATORCE",
            15: "QUINCE",
            16: "DIECISEIS",
            17: "DIECISIETE",
            18: "DIECIOCHO",
            19: "DIECINUEVE",
            20: "VEINTE",
            21: "VEINTIUNO",
            22: "VEINTIDOS",
            23: "VEINTITRES",
            24: "VEINTICUATRO",
            25: "VEINTICINCO",
            26: "VEINTISEIS",
            27: "VEINTISIETE",
            28: "VEINTIOCHO",
            29: "VEINTINUEVE"
        }

        decenas = [
            "",
            "",
            "VEINTE",
            "TREINTA",
            "CUARENTA",
            "CINCUENTA",
            "SESENTA",
            "SETENTA",
            "OCHENTA",
            "NOVENTA"
        ]

        centenas = [
            "",
            "CIENTO",
            "DOSCIENTOS",
            "TRESCIENTOS",
            "CUATROCIENTOS",
            "QUINIENTOS",
            "SEISCIENTOS",
            "SETECIENTOS",
            "OCHOCIENTOS",
            "NOVECIENTOS"
        ]

        def convertir_menor_1000(n):

            if n == 0:
                return ""

            if n == 100:
                return "CIEN"

            texto = ""

            if n >= 100:

                texto += centenas[n // 100]

                n %= 100

                if n:
                    texto += " "

            if 10 <= n <= 29:

                texto += especiales[n]

                return texto.strip()

            if n >= 30:

                texto += decenas[n // 10]

                n %= 10

                if n:
                    texto += " Y "

            if 0 < n < 10:
                texto += unidades[n]

            return texto.strip()

        numero = round(float(numero), 2)

        enteros = int(numero)

        centavos = int(
            round(
                (numero - enteros) * 100
            )
        )

        partes = []

        millones = enteros // 1_000_000

        resto = enteros % 1_000_000

        miles = resto // 1000

        unidades_finales = resto % 1000

        if millones:

            if millones == 1:
                partes.append("UN MILLON")
            else:
                partes.append(
                    convertir_menor_1000(
                        millones
                    )
                    + " MILLONES"
                )

        if miles:

            if miles == 1:
                partes.append("MIL")
            else:
                partes.append(
                    convertir_menor_1000(
                        miles
                    )
                    + " MIL"
                )

        if unidades_finales:

            partes.append(
                convertir_menor_1000(
                    unidades_finales
                )
            )

        if not partes:
            partes.append("CERO")

        texto = " ".join(partes)

        return (
            f"{texto} PESOS "
            f"{centavos:02d}/100 M.N."
        )

    @staticmethod
    def eliminar_referencias_externas(libro):

        for hoja in libro.worksheets:

            for fila in hoja.iter_rows():

                for celda in fila:

                    valor = celda.value

                    if (
                        isinstance(valor, str)
                        and valor.startswith("=")
                        and "[" in valor
                        and "]" in valor
                    ):

                        celda.value = None